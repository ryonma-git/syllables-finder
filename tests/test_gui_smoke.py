import json
import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
import app


class GuiSmokeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="syllables-gui-test-")
        self.state = Path(self.temp.name) / "state.json"
        self.state_patch = patch.object(app, "STATE_PATH", str(self.state))
        self.state_patch.start()
        self.root = tk.Tk()
        self.root.withdraw()
        self.ui = app.VowelMarkerApp(self.root)
        self.root.update_idletasks()

    def tearDown(self):
        self.root.destroy()
        self.state_patch.stop()
        self.temp.cleanup()

    def put(self, text):
        self.ui.input_text.delete("1.0", "end")
        self.ui.input_text.insert("1.0", text)

    def test_preview(self):
        self.put("make a dream\n日本語 123")
        self.ui.on_preview()
        self.assertEqual(self.ui.preview.get("1.0", "end-1c"), "m[a]ke [a] dr[ea]m\n日本語 123")

    def test_library_filter_and_load(self):
        self.ui.lib_filter.set("twinkle")
        self.assertEqual(self.ui.song_list.size(), 1)
        self.ui.song_list.selection_set(0)
        self.ui.on_load_song()
        self.assertTrue(self.ui.input_text.get("1.0", "end-1c").startswith("Twinkle, twinkle, little star"))
        self.assertEqual(self.ui.sheet_title.get(), "Twinkle, Twinkle, Little Star")
        self.ui.lib_filter.set("no-such-synthetic-title")
        self.assertEqual(self.ui.song_list.size(), 0)

    def test_invalid_size_blocks_save(self):
        self.ui.font_size.set("0")
        with patch.object(app.messagebox, "showerror") as error, patch.object(app.filedialog, "asksaveasfilename") as dialog:
            self.ui.on_generate()
        error.assert_called_once()
        dialog.assert_not_called()

    def test_empty_input_blocks_save(self):
        self.put("  \n")
        with patch.object(app.messagebox, "showwarning") as warning, patch.object(app.filedialog, "asksaveasfilename") as dialog:
            self.ui.on_generate()
        warning.assert_called_once()
        dialog.assert_not_called()

    def test_generate_document(self):
        output = Path(self.temp.name) / "gui-output.docx"
        self.put("make\n日本語")
        with patch.object(app.filedialog, "asksaveasfilename", return_value=str(output)), patch.object(app.messagebox, "showinfo") as info, patch.object(app.messagebox, "showerror") as error:
            self.ui.on_generate()
        error.assert_not_called()
        info.assert_called_once()
        self.assertEqual([p.text for p in Document(output).paragraphs], ["make", "日本語"])

    def test_generate_pdf(self):
        output = Path(self.temp.name) / "gui-output.pdf"
        self.put("make\n日本語")
        self.ui.sheet_title.set("Synthetic Title")
        with patch.object(app.filedialog, "asksaveasfilename", return_value=str(output)) as dialog, patch.object(app.messagebox, "showinfo") as info, patch.object(app.messagebox, "showerror") as error:
            self.ui.on_generate_pdf()
        error.assert_not_called()
        info.assert_called_once()
        self.assertEqual(dialog.call_args.kwargs["initialfile"], "red_vowel_output.pdf")
        self.assertTrue(output.read_bytes().startswith(b"%PDF"))

    def test_generate_document_uses_title(self):
        output = Path(self.temp.name) / "titled.docx"
        self.ui.sheet_title.set("Synthetic Title")
        with patch.object(app.filedialog, "asksaveasfilename", return_value=str(output)), patch.object(app.messagebox, "showinfo"):
            self.ui.on_generate()
        header = Document(output).sections[0].header
        self.assertIn("Synthetic Title", "".join(c.text for t in header.tables for r in t.rows for c in r.cells))

    def test_cancel_does_not_report_success(self):
        with patch.object(app.filedialog, "asksaveasfilename", return_value=""), patch.object(app.messagebox, "showinfo") as info:
            self.ui.on_generate()
        info.assert_not_called()

    def test_notice_state_is_local(self):
        self.assertTrue(self.ui._should_show_notice())
        self.ui.on_dismiss_notice()
        self.assertFalse(self.ui._should_show_notice())
        self.assertIn("notice_dismissed_at", json.loads(self.state.read_text()))


if __name__ == "__main__":
    unittest.main()
