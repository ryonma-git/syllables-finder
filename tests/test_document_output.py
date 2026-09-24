import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx_writer import write_docx
from vowel_marker import preview_text


class DocumentOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="syllables-test-")
        self.output = Path(self.temp.name) / "synthetic.docx"

    def tearDown(self):
        self.temp.cleanup()

    def test_japanese_blank_lines_and_punctuation_survive(self):
        text = "Make a dream!\n\n日本語 123 (OK)\n"
        write_docx(text, str(self.output))
        doc = Document(self.output)
        self.assertEqual([p.text for p in doc.paragraphs], text.split("\n"))

    def test_red_nucleus_silent_e_and_font(self):
        write_docx("make", str(self.output), font_name="Arial", font_size=24)
        runs = Document(self.output).paragraphs[0].runs
        self.assertEqual([(r.text, str(r.font.color.rgb)) for r in runs],
                         [("m", "000000"), ("a", "FF0000"), ("ke", "000000")])
        self.assertTrue(all(r.font.name == "Arial" and r.font.size.pt == 24 for r in runs))

    def test_custom_font_and_size(self):
        write_docx("bloom", str(self.output), font_name="Helvetica", font_size=18)
        runs = Document(self.output).paragraphs[0].runs
        self.assertTrue(all(r.font.name == "Helvetica" and r.font.size.pt == 18 for r in runs))
        self.assertEqual("".join(r.text for r in runs if str(r.font.color.rgb) == "FF0000"), "oo")

    def test_options_and_syllable_marks(self):
        self.assertEqual(preview_text("make", use_exceptions=False, silent_e=False), "m[a]k[e]")
        self.assertEqual(preview_text("sky", use_exceptions=False, y_as_vowel=False), "sky")
        write_docx("ma•ke\n日本語", str(self.output))
        self.assertEqual([p.text for p in Document(self.output).paragraphs], ["make", "日本語"])


if __name__ == "__main__":
    unittest.main()
