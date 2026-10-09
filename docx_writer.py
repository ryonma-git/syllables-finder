# -*- coding: utf-8 -*-
"""
docx_writer.py

mark_line() の判定結果をもとに、母音核だけ赤字にした .docx を書き出す。

- 各行を1段落として追加（空行も保持）。本文の段落は入力の行だけで、先生が Word で
  そのまま赤字を直せる
- 1段落の中に黒字 run と赤字 run を混在させる
- 赤字 = RGBColor(255, 0, 0) / 黒字 = RGBColor(0, 0, 0)
- フォント・サイズは指定可能（デフォルト Arial / 24pt）。日本語は游ゴシック
- 行間は読みやすさのため 1.3、段落後に少し余白

デザイン（sheet_design.py。PDF 版 pdf_writer.py と共通）:
- A4 縦。各ページのヘッダーに紺の帯（タイトル）と、凡例・なまえ欄
- 空行で区切られた連ごとに、本文の左へ淡い縦線
- フッターに「n / 総ページ数」
"""

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

import sheet_design as D
from vowel_marker import mark_line

RED = RGBColor.from_string(D.RED)
BLACK = RGBColor.from_string(D.BLACK)

# OOXML の子要素の順序（Word は順序違いを破損扱いすることがあるため、挿入位置を守る）
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_TCPR_AFTER_BORDERS = (
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
    "w:cellMerge", "w:tcPrChange",
)
_TCPR_AFTER_SHD = _TCPR_AFTER_BORDERS[1:]
_TCBORDERS_ORDER = ("top", "left", "bottom", "right", "insideH", "insideV", "tl2br", "tr2bl")
_TCPR_AFTER_MAR = _TCPR_AFTER_BORDERS[3:]
_RPR_AFTER_SPACING = (
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
    "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl",
    "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
)


# ------------------------------------------------------------ run（文字）
def _set_east_asia_font(rpr_owner, name: str):
    """run / style の rPr に日本語（East Asian）フォントを設定する。"""
    rfonts = rpr_owner.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def _apply_run_font(run, font_name: str, size_pt: float, red: bool):
    run.font.name = font_name           # 欧文(ascii/hAnsi)
    _set_east_asia_font(run._r, D.JAPANESE_FONT)
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RED if red else BLACK


def _char_spacing(run, spacing_pt: float):
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(int(round(spacing_pt * 20))))
    run._r.get_or_add_rPr().insert_element_before(spacing, *_RPR_AFTER_SPACING)


def _styled_run(paragraph, text, size_pt, color, bold=False,
                font=D.UI_LATIN_FONT, spacing_pt=None, symbol=False):
    """ヘッダー・フッター用の装飾 run。symbol=True は ♪ ● など記号を日本語フォントで描く。"""
    run = paragraph.add_run(text)
    run.font.name = D.JAPANESE_FONT if symbol else font
    _set_east_asia_font(run._r, D.JAPANESE_FONT)
    if symbol:
        run._r.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    if spacing_pt:
        _char_spacing(run, spacing_pt)
    return run


def _tight(paragraph, space_after_pt=0.0, style=None):
    if style is not None:
        paragraph.style = style
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after_pt)
    pf.line_spacing = 1.0


# ------------------------------------------------------------ 罫線・塗り
def _border_el(tag, color, size_eighths, space_pt=0):
    el = OxmlElement(tag)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_eighths))       # 1/8 pt 単位
    el.set(qn("w:space"), str(int(space_pt)))
    el.set(qn("w:color"), color)
    return el


def _paragraph_border(paragraph, edge, color, width_pt, space_pt):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *_PPR_AFTER_PBDR)
    pbdr.append(_border_el("w:" + edge, color, int(round(width_pt * 8)), space_pt))


def _cell_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, *_TCPR_AFTER_SHD)


def _cell_border(cell, edge, color, width_pt):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.insert_element_before(borders, *_TCPR_AFTER_BORDERS)
    el = _border_el("w:" + edge, color, int(round(width_pt * 8)))
    later = _TCBORDERS_ORDER[_TCBORDERS_ORDER.index(edge) + 1:]
    for child in borders:
        if child.tag in [qn("w:" + name) for name in later]:
            child.addprevious(el)
            return
    borders.append(el)


def _cell_margins(cell, left_mm=0.0, right_mm=0.0):
    mar = OxmlElement("w:tcMar")
    for edge, mm in (("top", 0), ("left", left_mm), ("bottom", 0), ("right", right_mm)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), _twips(mm))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    cell._tc.get_or_add_tcPr().insert_element_before(mar, *_TCPR_AFTER_MAR)


def _twips(mm):
    return str(int(round(mm * 1440 / 25.4)))


def _fixed_table_layout(table, widths_mm):
    """表の幅と列幅を固定（ソフトによって tblW / tcW / gridCol のどれを見るかが違うので全部そろえる）。"""
    table.autofit = False                      # w:tblLayout type="fixed"
    tblw = table._tbl.tblPr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), _twips(sum(widths_mm)))
    for gridcol, mm in zip(table._tbl.tblGrid.gridCol_lst, widths_mm):
        gridcol.w = Mm(mm)
    for row in table.rows:
        for cell, mm in zip(row.cells, widths_mm):
            cell.width = Mm(mm)


def _mark_size(paragraph, size_pt):
    """段落記号（文字の無い段落の高さ）の文字サイズ。"""
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(round(size_pt * 2))))
    rpr.append(sz)
    paragraph._p.get_or_add_pPr().insert_element_before(rpr, "w:sectPr", "w:pPrChange")


def _field(paragraph, instr, size_pt, color):
    """PAGE / NUMPAGES などの差し込みフィールド。"""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " {} ".format(instr))
    paragraph._p.append(fld)
    run = _styled_run(paragraph, "1", size_pt, color)
    fld.append(run._r)                # 段落直下から fldSimple の中へ移す


# ------------------------------------------------------------ ヘッダー・フッター
def _build_header(doc, section, title):
    header = section.header
    header.is_linked_to_previous = False
    style = doc.styles["Header"]
    name_w = D.NAME_FIELD_WIDTH_MM
    notes_w = D.NOTES_WIDTH_MM
    # 3列の格子: 帯は [タイトル(0+1列) | 音符(2列)]、次に黄色の線（全列）、
    # その下は [凡例(0列) | なまえ(1+2列)]
    widths = (D.TEXT_WIDTH_MM - name_w, name_w - notes_w, notes_w)

    table = header.add_table(rows=3, cols=3, width=Mm(D.TEXT_WIDTH_MM))
    _fixed_table_layout(table, widths)

    # ヘッダーは段落で終わる必要があるので、既定の空段落を表の後ろへ回して極小にする
    tail = header.paragraphs[0]
    header._element.remove(tail._p)
    header._element.append(tail._p)
    _tight(tail, style=style)
    tail.paragraph_format.line_spacing = Pt(1)
    _mark_size(tail, 1)

    # 1行目: 紺の帯
    band_row = table.rows[0]
    band_row.height = Mm(D.BAND_HEIGHT_MM)
    band_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    band = table.cell(0, 0).merge(table.cell(0, 1))
    notes = table.cell(0, 2)
    for cell in (band, notes):
        _cell_shading(cell, D.NAVY)
    # セルの境目に白い筋が出るソフトがあるので、帯と同じ色の線でふさぐ
    _cell_border(band, "right", D.NAVY, 0.5)
    _cell_border(notes, "left", D.NAVY, 0.5)
    for cell in (band, notes):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_margins(band, D.BAND_PAD_X_MM, 2.0)
    _cell_margins(notes, 0, D.BAND_PAD_X_MM)

    brand = band.paragraphs[0]
    _tight(brand, space_after_pt=1, style=style)
    _styled_run(brand, D.BRAND_MARK, D.BRAND_SIZE_PT, D.SKY, symbol=True)
    _styled_run(brand, " " + D.BRAND_TEXT, D.BRAND_SIZE_PT, D.SKY, bold=True,
                spacing_pt=D.BRAND_CHAR_SPACING_PT)

    heading = band.add_paragraph()
    _tight(heading, style=style)
    shown = D.display_title(title)
    _styled_run(heading, shown, D.title_size_pt(shown), D.WHITE, bold=True)

    mark = notes.paragraphs[0]
    _tight(mark, style=style)
    mark.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    mark.paragraph_format.line_spacing = Pt(D.NOTES_SIZE_PT * 1.1)   # 固定（帯が伸びないように）
    _styled_run(mark, D.NOTES_MARK, D.NOTES_SIZE_PT, D.NOTES, symbol=True)

    # 2行目: 黄色のアクセント線（1セルの塗りなので、上の帯の境目の線と交わらない）
    accent_row = table.rows[1]
    accent_row.height = Pt(D.ACCENT_WIDTH_PT)
    accent_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    accent = table.cell(1, 0).merge(table.cell(1, 2))
    _cell_shading(accent, D.ACCENT)
    _cell_margins(accent, 0, 0)
    stripe = accent.paragraphs[0]
    _tight(stripe, style=style)
    stripe.paragraph_format.line_spacing = Pt(1)
    _mark_size(stripe, 1)

    # 3行目: 凡例（左）・なまえ欄（右）
    info_row = table.rows[2]
    info_row.height = Mm(D.INFO_HEIGHT_MM)
    info_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    legend_cell = table.cell(2, 0)
    name_cell = table.cell(2, 1).merge(table.cell(2, 2))
    for cell in (legend_cell, name_cell):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_margins(legend_cell, 1.0, 0)
    _cell_margins(name_cell, 0, 0)

    legend = legend_cell.paragraphs[0]
    _tight(legend, style=style)
    _styled_run(legend, "●", D.LEGEND_DOT_SIZE_PT, D.RED, symbol=True)
    _styled_run(legend, " " + D.LEGEND_TEXT, D.INFO_SIZE_PT, D.NAVY, bold=True)
    _styled_run(legend, "　　" + D.LEGEND_EXAMPLE_LABEL + " ", D.INFO_SIZE_PT, D.MUTED)
    for fragment, red in D.LEGEND_EXAMPLE:
        _styled_run(legend, fragment, D.INFO_SIZE_PT + 1.5,
                    D.RED if red else D.BLACK, bold=True)

    # なまえ欄: ラベルの下に、欄の幅いっぱいの記入線（段落の下罫線）
    name = name_cell.paragraphs[0]
    _tight(name, style=style)
    _paragraph_border(name, "bottom", D.MUTED, 0.75, 2)
    _styled_run(name, D.NAME_LABEL, D.INFO_SIZE_PT, D.NAVY, bold=True)


def _build_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    _tight(p)
    _paragraph_border(p, "top", D.RULE, 0.75, 4)
    p.paragraph_format.tab_stops.add_tab_stop(
        Mm(D.TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
    _styled_run(p, D.FOOTER_TEXT, D.FOOTER_SIZE_PT, D.MUTED,
                spacing_pt=0.4)
    _styled_run(p, "\t", D.FOOTER_SIZE_PT, D.MUTED)
    _field(p, "PAGE", D.FOOTER_SIZE_PT, D.NAVY)
    _styled_run(p, " / ", D.FOOTER_SIZE_PT, D.MUTED)
    _field(p, "NUMPAGES", D.FOOTER_SIZE_PT, D.MUTED)


def _setup_header_footer_styles(doc):
    """Header/Footer スタイルの既定タブ（Letter 用の中央・右）を A4 用に直し、文字を小さくする。
    表セル内の段落記号やページ番号フィールドの大きさもここに従う。"""
    for name in ("Header", "Footer"):
        style = doc.styles[name]
        tabs = style.paragraph_format.tab_stops
        tabs.clear_all()
        tabs.add_tab_stop(Mm(D.TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
        style.font.name = D.UI_LATIN_FONT
        style.font.size = Pt(D.FOOTER_SIZE_PT)
        style.font.color.rgb = RGBColor.from_string(D.MUTED)
        _set_east_asia_font(style.element, D.JAPANESE_FONT)


def _setup_page(section):
    section.page_width = Mm(D.PAGE_WIDTH_MM)
    section.page_height = Mm(D.PAGE_HEIGHT_MM)
    section.left_margin = Mm(D.MARGIN_LEFT_MM)
    section.right_margin = Mm(D.MARGIN_RIGHT_MM)
    section.top_margin = Mm(D.MARGIN_TOP_MM)
    section.bottom_margin = Mm(D.MARGIN_BOTTOM_MM)
    section.header_distance = Mm(D.HEADER_DISTANCE_MM)
    section.footer_distance = Mm(D.FOOTER_DISTANCE_MM)


# ------------------------------------------------------------ 本体
def write_docx(text: str,
               filepath: str,
               font_name: str = "Arial",
               font_size: float = 24,
               use_exceptions: bool = True,
               silent_e: bool = True,
               y_as_vowel: bool = True,
               title: str = ""):
    """text を処理して filepath に .docx を保存する。title はヘッダーの帯に出す（空なら既定の見出し）。"""
    doc = Document()
    doc.core_properties.title = D.display_title(title)

    # 既定スタイル（Normal）にもフォントを設定しておく（空行・日本語のフォールバック用）
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(font_size)
    _set_east_asia_font(normal.element, D.JAPANESE_FONT)

    section = doc.sections[0]
    _setup_page(section)
    _setup_header_footer_styles(doc)
    _build_header(doc, section, title)
    _build_footer(section)

    opts = dict(
        use_exceptions=use_exceptions,
        silent_e=silent_e,
        y_as_vowel=y_as_vowel,
    )

    bar_space = D.STANZA_BAR_GAP_PT
    for line in text.split("\n"):
        fragments = mark_line(line, **opts)
        paragraph = doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.line_spacing = D.LINE_SPACING    # 教材として読みやすい行間
        pf.space_after = Pt(D.SPACE_AFTER_PT)   # 段落間を詰めすぎない
        pf.left_indent = Mm(D.BODY_INDENT_MM)

        # 文字のある行だけ左に縦線。空行で途切れるので、連ごとの線になる
        if "".join(f for f, _ in fragments).strip():
            _paragraph_border(paragraph, "left", D.STANZA_BAR,
                              D.STANZA_BAR_WIDTH_PT, bar_space)

        for fragment, red in fragments:
            run = paragraph.add_run(fragment)
            _apply_run_font(run, font_name, font_size, red)

    doc.save(filepath)
    return filepath
